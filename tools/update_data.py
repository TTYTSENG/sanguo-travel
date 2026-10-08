"""Export only public source sheets. Never export the personal journal."""
import argparse, json, re, gzip
from pathlib import Path
import openpyxl

def export(source, output):
    w = openpyxl.load_workbook(source, read_only=True, data_only=True)
    def rows(name, width):
        return [[v if v is not None else '' for v in r] for r in
                w[name].iter_rows(min_row=5, max_col=width, values_only=True) if r[0]]
    regions = rows('地區對照', 9)
    data = dict(cities=[r[:6] for r in regions],
                aliases=[[r[7], r[8]] for r in w['地區對照'].iter_rows(min_row=5,max_col=9,values_only=True) if r[7] and r[8]],
                maps=rows('古今地名對照',6), museum=rows('博物館文物',10),
                master=rows('史料與小說資料',14), relations=rows('關聯索引',9),
                stations=rows('12306車站庫',7), examples=[])
    w.close()
    ids={r[0] for r in data['cities']}
    if len(ids)!=len(data['cities']): raise ValueError('Duplicate city IDs')
    if len({r[0] for r in data['master']})!=len(data['master']): raise ValueError('Duplicate source IDs')
    kept=[]
    for r in data['relations']:
        index=int(r[5])-5
        if r[1] not in ids or index<0 or index>=len(data['master']) or data['master'][index][0]!=r[6]:
            raise ValueError('Invalid relationship or source row: '+str(r[0]))
        if r[1]=='C001' and not re.search(r'成都(?!邑)',data['master'][index][10]): continue
        kept.append(r)
    data['relations']=kept
    p=Path(output)
    # English descriptions use the same city and museum ordering as the workbook.
    if p.exists():
        old=json.loads(p.read_text(encoding='utf-8').split('=',1)[1].rstrip(';\n'))
        if [r[0] for r in old['cities']]!=[r[0] for r in data['cities']] or [r[:4] for r in old['museum']]!=[r[:4] for r in data['museum']]:
            raise ValueError('City/museum catalog changed. Update docs/i18n.js ordering and translations before export.')
    p.write_text('globalThis.SANGUO_DATA='+json.dumps(data,ensure_ascii=False,separators=(',',':'))+';',encoding='utf-8')
    p.with_suffix('.json.gz').write_bytes(gzip.compress(json.dumps(data,ensure_ascii=False,separators=(',',':')).encode('utf-8'),mtime=0))
    print({k:len(v) for k,v in data.items()})

if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--xlsx',required=True)
    p.add_argument('--output',default='docs/data.js')
    a=p.parse_args();export(a.xlsx,a.output)
